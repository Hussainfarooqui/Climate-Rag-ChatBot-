from pypdf import PdfReader
import docx
from pptx import Presentation
from bs4 import BeautifulSoup
import pandas as pd
from .utils import save_uploaded_file
from pathlib import Path

def load_pdf(path):
    text_chunks = []
    reader = PdfReader(path)
    for page in reader.pages:
        text_chunks.append(page.extract_text() or "")
    return "\n".join(text_chunks)

def load_docx(path):
    doc = docx.Document(path)
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

def load_pptx(path):
    prs = Presentation(path)
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
    return "\n".join(texts)

def load_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def load_csv(path):
    try:
        df = pd.read_csv(path)
        return df.to_string(index=False)
    except:
        return load_txt(path)

def load_html(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f, "html.parser")
        return soup.get_text(separator="\n")

def load_file(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        path = save_uploaded_file(uploaded_file, suffix=".pdf")
        return load_pdf(path)
    elif name.endswith(".docx"):
        path = save_uploaded_file(uploaded_file, suffix=".docx")
        return load_docx(path)
    elif name.endswith(".pptx"):
        path = save_uploaded_file(uploaded_file, suffix=".pptx")
        return load_pptx(path)
    elif name.endswith(".txt") or name.endswith(".md"):
        path = save_uploaded_file(uploaded_file, suffix=".txt")
        return load_txt(path)
    elif name.endswith((".csv",".tsv",".xls",".xlsx")):
        path = save_uploaded_file(uploaded_file, suffix=Path(uploaded_file.name).suffix)
        return load_csv(path)
    elif name.endswith((".html",".htm")):
        path = save_uploaded_file(uploaded_file, suffix=".html")
        return load_html(path)
    else:
        try:
            return uploaded_file.getvalue().decode("utf-8", errors="ignore")
        except:
            return ""
